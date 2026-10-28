# -*- coding: utf-8 -*-
"""用 python-pptx 验证 BiliPptx builder 生成的 PPTX 是否结构合法、可正常打开。"""
import os
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "test-output.pptx")

# 1) zip 结构完整性
import xml.etree.ElementTree as ET

with zipfile.ZipFile(PPTX) as z:
    bad = z.testzip()
    assert bad is None, f"zip 损坏: {bad}"
    names = z.namelist()
    print(f"zip OK，共 {len(names)} 个部件")
    for must in (
        "[Content_Types].xml",
        "_rels/.rels",
        "ppt/presentation.xml",
        "ppt/slideMasters/slideMaster1.xml",
        "ppt/slideLayouts/slideLayout1.xml",
        "ppt/theme/theme1.xml",
        "ppt/slides/slide4.xml",
        "ppt/media/image4.png",
    ):
        assert must in names, f"缺少部件: {must}"
    # 所有 XML 部件均可解析（well-formed）
    xml_parts = [n for n in names if n.endswith(".xml") or n.endswith(".rels")]
    for n in xml_parts:
        ET.fromstring(z.read(n))
    print(f"XML well-formed OK，共 {len(xml_parts)} 个 XML 部件")

# 2) python-pptx 能打开并解析
prs = Presentation(PPTX)
assert len(prs.slides) == 4, f"应为 4 页，实际 {len(prs.slides)}"
assert prs.slide_width == Emu(12192000), f"宽不符: {prs.slide_width}"
assert prs.slide_height == Emu(6858000), f"高不符: {prs.slide_height}"

for i, slide in enumerate(prs.slides, 1):
    shapes = list(slide.shapes)
    assert len(shapes) == 1, f"第 {i} 页应有 1 个形状，实际 {len(shapes)}"
    pic = shapes[0]
    assert pic.shape_type is not None, f"第 {i} 页形状类型缺失"
    img = pic.image
    assert img.blob and len(img.blob) > 1000, f"第 {i} 页图片数据异常"
    assert img.ext == "png", f"第 {i} 页图片扩展名: {img.ext}"
    print(f"slide {i}: OK  图片 {len(img.blob)} bytes, {img.size}")

print("\npython-pptx 验证全部通过 ✓")
